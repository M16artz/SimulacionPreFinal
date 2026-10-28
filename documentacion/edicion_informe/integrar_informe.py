"""Integra las secciones sugeridas en el artículo conservando sus autores y referencias."""
from copy import deepcopy
from pathlib import Path
import re
import shutil
import sys
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

RAIZ = Path(__file__).resolve().parents[2]
CARPETA = Path(__file__).resolve().parent
ORIGINAL = RAIZ / "InformeFinalSimulacion(2).docx"
RESPALDO = CARPETA / "InformeFinalSimulacion(2)_original.docx"
SALIDA = CARPETA / "InformeFinalSimulacion(2)_integrado.docx"
if not RESPALDO.exists():
    shutil.copy2(ORIGINAL, RESPALDO)
doc = Document(RESPALDO)
for elemento in doc._element.iter():
    valor = elemento.get(qn("w:w"))
    if valor and re.fullmatch(r"\d+\.\d+", valor):
        elemento.set(qn("w:w"), str(round(float(valor))))

referencias = [deepcopy(p._p) for p in doc.paragraphs if re.match(r"^\[\d+\]", p.text)]
variables = [[c.text for c in fila.cells] for fila in doc.tables[2].rows]
modelo_titulo = deepcopy(doc.paragraphs[7]._p.pPr)
inicio = next(p._p for p in doc.paragraphs if p.text == "D. Procedimiento")
cuerpo = doc._element.body
eliminar = False
for elemento in list(cuerpo):
    if elemento is inicio:
        eliminar = True
    if eliminar and elemento.tag != qn("w:sectPr"):
        cuerpo.remove(elemento)

def texto_limpio(texto):
    texto = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", texto)
    return texto.replace("`", "").replace("**", "")

def formato_parrafo(p, tamano=10, sangria=True):
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(4)
    pf.line_spacing = 1.0
    pf.first_line_indent = Inches(0.14) if sangria else Pt(0)
    pf.widow_control = True
    pf.keep_with_next = False
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for r in p.runs:
        r.font.name = "Times New Roman"
        r.font.size = Pt(tamano)
        r.font.color.rgb = RGBColor(0, 0, 0)
        r.font.highlight_color = None

def poner_texto(p, texto, tamano=10):
    p.clear()
    for i, pedazo in enumerate(re.split(r"(\*\*.*?\*\*)", texto)):
        if not pedazo:
            continue
        r = p.add_run(texto_limpio(pedazo))
        r.bold = pedazo.startswith("**")
    formato_parrafo(p, tamano)
    return p

def parrafo(texto, tamano=10):
    return poner_texto(doc.add_paragraph(style="normal"), texto, tamano)

def encabezado(texto, nivel):
    p = doc.add_paragraph(texto, style=f"Heading {nivel}")
    if nivel == 1:
        p._p.remove(p._p.pPr)
        p._p.insert(0, deepcopy(modelo_titulo))
    formato_parrafo(p, 10, False)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if nivel == 1 else WD_ALIGN_PARAGRAPH.LEFT
    for r in p.runs:
        r.italic = nivel > 1
        r.bold = nivel == 1
    return p

def ecuacion(texto):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(5)
    o = OxmlElement("m:oMath")
    r = OxmlElement("m:r")
    t = OxmlElement("m:t")
    t.text = texto
    r.append(t)
    o.append(r)
    p._p.append(o)

def formulacion():
    ecuacion("λ(t) = 20 grupos/h, si 12 ≤ t < 14 o 18 ≤ t < 20")
    ecuacion("λ(t) = 8 grupos/h, en las demás horas de admisión")
    parrafo("El tiempo de observación T se midió en minutos, desde la apertura hasta finalizar los servicios admitidos. Para cada recurso r, cᵣ representó su capacidad y T ocupado la suma de los tiempos de actividad de todos sus servidores.")
    ecuacion("Uᵣ = [T ocupado,r / (cᵣ T)] × 100")
    ecuacion("W̄q,i = (1/nᵢ) ∑ Wq,i,k")
    parrafo("En la expresión anterior, la suma abarcó los nᵢ registros de espera de la estación i. La longitud media de la cola se calculó integrando la cantidad Qᵢ(t) de grupos pendientes; para una trayectoria por eventos, la integral equivale a sumar Qᵢ por la duración de cada intervalo.")
    ecuacion("L̄q,i = (1/T) ∫₀ᵀ Qᵢ(t) dt")
    ecuacion("Wtotal,k = tsalida,k − tllegada,k")
    ecuacion("Paband = (Naband / Nllegadas) × 100")
    parrafo("La permanencia media consideró solamente los grupos que completaron el servicio. El porcentaje de abandono utilizó todos los grupos llegados como denominador. Para contrastar la ley de Little se emplearon grupos por minuto, minutos y número medio de grupos, con población y ventana compatibles [1].")
    ecuacion("L = λef W")
    parrafo("La longitud media de cola y el contraste de Little pertenecen a la ampliación 1.1.0; las comparaciones de los experimentos históricos se limitaron a los indicadores efectivamente registrados en la versión 1.0.0.")

ancho = doc.sections[-1].page_width - doc.sections[-1].left_margin - doc.sections[-1].right_margin
columnas = 2

def establecer_columnas(numero):
    global columnas
    if columnas == numero:
        return
    seccion = doc.add_section(WD_SECTION_START.CONTINUOUS)
    cols = seccion._sectPr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        seccion._sectPr.append(cols)
    cols.clear()
    cols.set(qn("w:num"), str(numero))
    cols.set(qn("w:equalWidth"), "1")
    cols.set(qn("w:space"), "360")
    columnas = numero

def romano(n):
    salida = ""
    for valor, simbolo in [(10, "X"), (9, "IX"), (5, "V"), (4, "IV"), (1, "I")]:
        while n >= valor:
            salida += simbolo
            n -= valor
    return salida

def cambiar_tablas(texto):
    return re.sub(r"([Tt]abla) (\d+)", lambda m: f"{m[1]} {romano(int(m[2]) + (int(m[2]) >= 2))}", texto)

def tabla(filas, titulo):
    establecer_columnas(1)
    p = parrafo(titulo, 9)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.keep_with_next = True
    t = doc.add_table(rows=0, cols=len(filas[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    n = len(filas[0])
    pesos = ([2.4, 4.6] if n == 2 else [1.2, 3.6, 2.2] if n == 3 else [1.15] + [(7.0 - 1.15) / (n - 1)] * (n - 1))
    for columna, peso in zip(t.columns, pesos):
        columna.width = int(ancho * peso / sum(pesos))
    for i, fila in enumerate(filas):
        row = t.add_row()
        pr = row._tr.get_or_add_trPr()
        pr.append(OxmlElement("w:cantSplit"))
        if i == 0:
            pr.append(OxmlElement("w:tblHeader"))
        for j, (celda, texto) in enumerate(zip(row.cells, fila)):
            celda.width = int(ancho * pesos[j] / sum(pesos))
            celda.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = poner_texto(celda.paragraphs[0], texto_limpio(texto), 9)
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if n <= 3 or j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.bold = i == 0
            tcpr = celda._tc.get_or_add_tcPr()
            bordes = OxmlElement("w:tcBorders")
            for borde in ("top", "left", "bottom", "right"):
                b = OxmlElement("w:" + borde)
                b.set(qn("w:val"), "single")
                b.set(qn("w:sz"), "4")
                b.set(qn("w:color"), "D9D9D9")
                bordes.append(b)
            tcpr.append(bordes)
            margenes = OxmlElement("w:tcMar")
            for lado in ("top", "left", "bottom", "right"):
                m = OxmlElement("w:" + lado)
                m.set(qn("w:w"), "80")
                m.set(qn("w:type"), "dxa")
                margenes.append(m)
            tcpr.append(margenes)
            if i == 0:
                shd = OxmlElement("w:shd")
                shd.set(qn("w:fill"), "E7EDF3")
                tcpr.append(shd)
    espaciador = parrafo("", 2)
    espaciador.paragraph_format.space_after = Pt(1)
    establecer_columnas(2)

# Ajustes puntuales de las partes conservadas del artículo.
for p in doc.paragraphs:
    if p.text.startswith("Resumen"):
        poner_texto(p, "**Resumen —** Se desarrolló un modelo de red de colas mediante simulación de eventos discretos en Python y SimPy para evaluar cantidades de mesas, meseros y cocineras en un restaurante. Se representaron llegadas variables por horario, asignación conjunta de mesas, abandono y atención con personal compartido. Se analizaron tres lotes por fuerza bruta, con cuatro o cinco réplicas por escenario y 183 combinaciones distintas. Los experimentos utilizaron datos sintéticos, llegadas pico de 20 grupos por hora y paciencia uniforme entre 0 y 20 minutos. En el lote de cinco réplicas y hasta tres meseros, la configuración 13/3/6 obtuvo la menor permanencia media observada, de 44,65 minutos, frente a 62,69 en la referencia 11/1/3. El abandono disminuyó de 31,45 % a 9,56 % y las personas atendidas aumentaron de 225,40 a 308,20 por jornada. Se identificaron rendimientos marginales reducidos y diferencias entre minimizar permanencia y la eficiencia de Pareto. Las conclusiones se limitan a las configuraciones y supuestos analizados, sin afirmar validación empírica ni un óptimo económico universal.", 9)
    elif p.text == "B.. Tipo de estudio":
        p.text = "B. Tipo de estudio"
    elif p.text.startswith("Para desarrollar el modelo se utilizará"):
        poner_texto(p, "Se utilizó Python y SimPy para implementar el motor de eventos discretos. NumPy gestionó la generación aleatoria y SciPy los intervalos estadísticos. Los resultados se conservaron en CSV y JSON. La aplicación incorporó una API FastAPI y una interfaz React para configurar escenarios, reproducir eventos y consultar gráficos. Las figuras del análisis se elaboraron con ReportLab Charts a partir de las exportaciones verificadas.")
    elif p.text.startswith("El modelo se fundamentará"):
        poner_texto(p, "El modelo se fundamentó en una red de colas con recursos compartidos. Las expresiones M/M/s se conservaron como referencia teórica, sin utilizarlas directamente para producir los resultados experimentales. Esta elección consideró los antecedentes de Dey et al. [2], Buzali et al. [8] y Samad et al. [9].")
    elif "un 40 % el tiempo promedio de permanencia" in p.text:
        poner_texto(p, p.text.replace("ocupadas.Respecto", "ocupadas. Respecto").replace("un 40 % el tiempo promedio de permanencia", "un 40 % el tiempo de espera"))

# Título y cabecera de autoría a todo el ancho antes del texto en dos columnas.
titulo = doc.paragraphs[0]
titulo.style = doc.styles["Title"]
titulo.paragraph_format.first_line_indent = Pt(0)
titulo.paragraph_format.space_after = Pt(10)
for r in titulo.runs:
    r.font.color.rgb = RGBColor(0, 0, 0)
    r.font.name = "Times New Roman"
    r.font.size = Pt(20)
    r.font.highlight_color = None
cabecera = doc.tables[0]
for i, columna in enumerate(cabecera.columns):
    columna.width = int(ancho / 4)
for celda in cabecera.rows[0].cells:
    celda.width = int(ancho / 4)
    for p in celda.paragraphs:
        formato_parrafo(p, 8.5, False)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
salto = next(p._p for p in doc.paragraphs if p._p.find(".//" + qn("w:sectPr")) is not None)
salto.getparent().remove(salto)
cabecera._tbl.addnext(salto)

# Texto de la propuesta, sin notas de edición ni anexo de trabajo.
md = (RAIZ / "Contenido_Informe_Secciones_6_4_a_10.md").read_text(encoding="utf-8")
md = md[md.index("## 6.4. Procedimiento"):md.index("## Fuentes y trazabilidad")].strip().removesuffix("---").strip()
lineas = md.splitlines()
i = 0
titulo_tabla = None
subdiscusion = 0
while i < len(lineas):
    linea = lineas[i].strip()
    if not linea or linea == "---":
        i += 1
        continue
    if linea.startswith("## "):
        nombre = linea[3:]
        if nombre.startswith("6.4."):
            encabezado("D. Procedimiento", 2)
        elif nombre.startswith("6.5."):
            encabezado("E. Evaluación", 2)
        else:
            encabezado(re.sub(r"^\d+\.\s*", "", nombre).upper(), 1)
        i += 1
        continue
    if linea.startswith("### "):
        nombre = linea[4:]
        if nombre.startswith("Etapa 6."):
            formulacion()
        if nombre.startswith("7.1."):
            nombre = "A. Resultados del objetivo específico 1"
        elif nombre.startswith("7.2."):
            nombre = "B. Resultados del objetivo específico 2"
        elif nombre.startswith("6.5."):
            nombre = re.sub(r"^6\.5\.(\d+)\.", r"\1.", nombre)
        elif nombre.startswith("8."):
            subdiscusion += 1
            nombre = re.sub(r"^8\.\d+\.", chr(64 + subdiscusion) + ".", nombre)
        encabezado(nombre, 3 if linea[4:].startswith(("Etapa", "6.5.")) else 2)
        i += 1
        continue
    if linea.startswith("#### "):
        nombre = re.sub(r"^7\.2\.(\d+)\.", r"\1.", linea[5:])
        encabezado(nombre, 3)
        i += 1
        continue
    if linea.startswith("**Tabla "):
        titulo_tabla = cambiar_tablas(texto_limpio(linea))
        i += 1
        continue
    if linea.startswith("|"):
        filas = []
        while i < len(lineas) and lineas[i].strip().startswith("|"):
            valores = [v.strip() for v in lineas[i].strip().strip("|").split("|")]
            if not all(re.fullmatch(r":?-+:?", v) for v in valores):
                filas.append(valores)
            i += 1
        tabla(filas, titulo_tabla)
        if titulo_tabla.startswith("Tabla I."):
            tabla(variables, "Tabla II. Variables del modelo")
        titulo_tabla = None
        continue
    if linea.startswith("!["):
        imagen = re.search(r"\]\(([^)]+)\)", linea)[1]
        nombre = Path(imagen).name
        ruta = CARPETA / nombre
        if not ruta.exists():
            ruta = RAIZ / imagen
        establecer_columnas(1)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(ruta), width=ancho)
        i += 1
        while i < len(lineas) and not lineas[i].strip():
            i += 1
        if i < len(lineas) and lineas[i].strip().startswith("**Figura"):
            cap = parrafo(lineas[i].strip(), 9)
            cap.paragraph_format.first_line_indent = Pt(0)
            i += 1
        establecer_columnas(2)
        continue
    if linea == r"\[":
        ecuacion = []
        i += 1
        while i < len(lineas) and lineas[i].strip() != r"\]":
            ecuacion.append(lineas[i])
            i += 1
        # Ecuaciones nativas y editables, sin dejar instrucciones LaTeX.
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        o = OxmlElement("m:oMath")
        r = OxmlElement("m:r")
        t = OxmlElement("m:t")
        t.text = "x̄ ± t₀,₉₇₅;ₙ₋₁ · s/√n" if "bar{x}" in "".join(ecuacion) else "(M, S, C, P̄aband, W̄mesa)"
        r.append(t)
        o.append(r)
        p._p.append(o)
        i += 1
        continue
    fragmentos = [linea]
    i += 1
    while i < len(lineas) and lineas[i].strip() and not lineas[i].strip().startswith(("#", "|", "![", "**Tabla", r"\[")):
        fragmentos.append(lineas[i].strip())
        i += 1
    parrafo(cambiar_tablas(" ".join(fragmentos)))

encabezado("REFERENCIAS", 1)
for elemento in referencias:
    cuerpo.insert(len(cuerpo) - 1, elemento)
    from docx.text.paragraph import Paragraph
    p = Paragraph(elemento, doc._body)
    formato_parrafo(p, 9, False)
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.first_line_indent = Inches(-0.18)
encabezado("DECLARACIÓN DE IA", 1)
parrafo("Se utilizó una herramienta de inteligencia artificial como apoyo para la organización del texto, la integración de secciones y el análisis de archivos experimentales. Las cifras se contrastaron con los CSV y sus registros JSON; la herramienta no sustituyó la simulación ni la validación empírica. La revisión y responsabilidad del contenido corresponden a los autores.")

for nombre in ("Heading 1", "Heading 2", "Heading 3", "Title"):
    doc.styles[nombre].font.color.rgb = RGBColor(0, 0, 0)
doc.core_properties.title = titulo.text
doc.save(SALIDA)
print(SALIDA)
